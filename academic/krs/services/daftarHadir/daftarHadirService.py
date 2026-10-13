from pathlib import Path
import io
import typing
from dataclasses import dataclass

from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.utils import column_index_from_string
from openpyxl.utils.cell import coordinate_to_tuple

from academic.models import MataKuliah, DetailKRS, KRS, Mahasiswa, Kelas


class DatakuType(typing.TypedDict):
    daftar_mahasiswa: list[Mahasiswa]
    kelas: Kelas
    krs: KRS
    # periode_krs: typing.Any


@dataclass(frozen=True)
class DatakuSchema:
    daftar_mahasiswa: list[Mahasiswa]
    kelas: Kelas
    krs: KRS


@dataclass(frozen=True)
class Result:
    is_success: bool
    errors: list[str, list[str]] | None = None
    data: any = None
    messages: str | None = None


class FieldDaftarHadirType(typing.TypedDict):
    MATA_KULIAH_SKS: str
    SEMESTER_RUANG: str
    DOSEN_PEMBINA: str
    NIDN: str
    WAKTU: str
    NAMA_KELAS: str
    NIM: str | list[str]
    NAMA: str | list[str]
    DOSEN_PEMBINA_NAMA: str
    DOSEN_PEMBINA_NIDN: str


class FieldJurnalType(typing.TypedDict):
    TAHUN_AKADEMIK: str
    MATA_KULIAH_SKS: str
    SEMESTER: str
    DOSEN_PEMBINA: str


fieldSet_daftar_hadir: FieldDaftarHadirType = {
    "MATA_KULIAH_SKS": "E8",
    "SEMESTER_RUANG": "E9",
    "DOSEN_PEMBINA": "E10",
    "NIDN": "L10",
    "WAKTU": "E11",
    "NAMA_KELAS": "S8",
    "NIM": "B17",  # cell pertama
    "NAMA": "C17",  # cell pertama
    "DOSEN_PEMBINA_NAMA": "O65",
    "DOSEN_PEMBINA_NIDN": "P66",
}

fieldSet_jurnal_perkuliahan: FieldJurnalType = {
    "TAHUN_AKADEMIK": "E8",
    "MATA_KULIAH_SKS": "E9",
    "SEMESTER": "E10",
    "DOSEN_PEMBINA": "E11",
}


class Query:
    @staticmethod
    def mydata(*, kelas_id: int) -> DatakuSchema:
        detail_krs = (
                DetailKRS.objects
                .filter(kelas_id=kelas_id)
                .select_related(
                    "krs__mahasiswa", "kelas__mata_kuliah"
                )
                .order_by("krs__mahasiswa__nim")
        )

        daftar_mahasiswa = [detail.krs.mahasiswa for detail in detail_krs]

        kelas = detail_krs[0].kelas
        krs = detail_krs[0].krs

        return DatakuSchema(daftar_mahasiswa=daftar_mahasiswa, kelas=kelas, krs=krs)

        # return {
        #     "daftar_mahasiswa": daftar_mahasiswa,
        #     "kelas": kelas,
        #     # "periode_krs": None,
        #     "krs": krs
        # }


def angka_ke_romawi(angka: int) -> str:
    romawi = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII"]
    return romawi[angka - 1]


class ExcelService:
    def __init__(self):
        pass

    @staticmethod
    def simpan_excel(wb) -> io.BytesIO:
        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        return output

    @staticmethod
    def sisipkan_array(
        ws, *, kolom: str, data: list, baris_awal: int = 1, coordinate=None
    ):
        if coordinate:
            baris, kolom = coordinate_to_tuple(coordinate=coordinate)

            for row, value in enumerate(data, start=baris):
                ws.cell(row=row, column=kolom, value=value)
        else:
            nomor_kolom = column_index_from_string(kolom)

            for row, value in enumerate(data, start=baris_awal):
                ws.cell(row=row, column=nomor_kolom, value=value)

    @staticmethod
    def sisipkan_nilai_by_coordinat(ws, *, coordinate: str, value: str):
        baris, kolom = coordinate_to_tuple(coordinate=coordinate)

        ws.cell(row=baris, column=kolom, value=value)

    @staticmethod
    def excecute(kelas_id: int, template_path: Path, **kwargs) -> Result:
        data = Query.mydata(kelas_id=kelas_id)
        daftar_mahasiswa = data.daftar_mahasiswa
        kelas = data.kelas
        krs = data.krs

        wb = load_workbook(template_path)
        ws_daftar_hadir = wb["KELAS"]
        ws_jurnal_perkuliahan = wb["JURNAL PERKULIAHAN"]

        dosen_nama = kwargs.get("DOSEN_PEMBINA", "")
        dosen_nidn = kwargs.get("NIDN", "")

        value_daftar_hadir: FieldDaftarHadirType = {
            "DOSEN_PEMBINA": dosen_nama,
            "NIDN": dosen_nidn,
            "DOSEN_PEMBINA_NIDN": dosen_nidn,
            "DOSEN_PEMBINA_NAMA": dosen_nama,
            "NAMA": [mahasiswa.nama for mahasiswa in daftar_mahasiswa],
            "NIM": [mahasiswa.nim for mahasiswa in daftar_mahasiswa],
            "WAKTU": kelas.waktu,
            "SEMESTER_RUANG": f"{angka_ke_romawi(kelas.mata_kuliah.semester)} / {kelas.ruangan or 'R207'}",
            "MATA_KULIAH_SKS": f"{kelas.mata_kuliah.nama} / {kelas.mata_kuliah.sks} SKS",
            "NAMA_KELAS": f"KELAS {kelas.nama.upper()}",
        }

        # 'GANJIL T.A. 2026/2027
        
        value_jurnal_perkuliahan: FieldJurnalType = {
            "DOSEN_PEMBINA": dosen_nama,
            "MATA_KULIAH_SKS": f"{kelas.mata_kuliah.nama.upper()} / {kelas.mata_kuliah.sks} SKS",
            "SEMESTER": angka_ke_romawi(kelas.mata_kuliah.semester),
            "TAHUN_AKADEMIK": f"{krs.semester.upper()} T.A {krs.tahun_ajaran}",
        }

        for key in fieldSet_daftar_hadir.keys():
            if key == "NAMA" or key == "NIM":
                continue
            ExcelService.sisipkan_nilai_by_coordinat(
                ws_daftar_hadir,
                coordinate=fieldSet_daftar_hadir[key],
                value=value_daftar_hadir[key],
            )
            
        for key in fieldSet_jurnal_perkuliahan.keys():
            ExcelService.sisipkan_nilai_by_coordinat(
                ws_jurnal_perkuliahan,
                coordinate=fieldSet_jurnal_perkuliahan[key],
                value=value_jurnal_perkuliahan[key],
            )

        # baris, kolom = coordinate_to_tuple(coordinate=fieldSet_daftar_hadir["NIM"])
        ExcelService.sisipkan_array(
            ws_daftar_hadir,
            kolom=fieldSet_daftar_hadir["NIM"],
            coordinate=fieldSet_daftar_hadir["NIM"],
            data=value_daftar_hadir["NIM"],
        )

        ExcelService.sisipkan_array(
            ws_daftar_hadir,
            kolom=fieldSet_daftar_hadir["NAMA"],
            coordinate=fieldSet_daftar_hadir["NAMA"],
            data=value_daftar_hadir["NAMA"],
        )

        # return ExcelService.simpan_excel(wb)
        return Result(is_success=True, data=ExcelService.simpan_excel(wb))

        # ExcelService.sisipkan_nilai_by_coordinat(ws, coordinate="E8", value=mata_kuliah.mata_kuliah)
        # ExcelService.sisipkan_nilai_by_coordinat(ws, coordinate="E9", value=mata_kuliah.semester)
        # ExcelService.sisipkan_nilai_by_coordinat(ws, coordinate="E10", value=mata_kuliah.dosen_pembina)
        # ExcelService.sisipkan_nilai_by_coordinat(ws, coordinate="E11", value=mata_kuliah.waktu)
